# A program used to analyze data prior to distribution.
# sammatime, 2026
import datetime
import google.generativeai as genai
import html
import logging
import openpyxl
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.properties import PageSetupProperties
import re
import ssl
from time import sleep
import traceback
from urllib.parse import quote_plus
from urllib.request import Request, urlopen
import xml.etree.ElementTree as ElementTree

# Config keys
EMAIL_CONFIG = "email_config"
ATTACHMENT = "attachment"

# SQL queries
SELECT_STOCK_NAME_AND_ACRONYM = "SELECT stock_name, acronym FROM STOCK WHERE stock_id={};"

# A price of -1 means the price is missing (-1 == -1.0, so both forms match)
MISSING_PRICE = -1

# AI settings
# Models to try in order: if one fails (busy with "503 high demand", or a "504 Deadline Exceeded"),
# wait and ask the next one: 30 seconds before the second, 60 before the third and 90 before the fourth
# (Flash-Lite last: a lighter model, often still free when the others are busy)
GEMINI_MODELS = ["gemini-3.8-flash", "gemini-3.7-flash", "gemini-3.6-flash", "gemini-3.1-flash-lite"]
AI_RETRY_WAITS = [30, 60, 90]
AI_ERROR_ANSWER = "An error occurred in querying the AI agent."
TOP_COUNT = 3
# Distribution runs once a day and nothing waits on it, so Gemini can take up to 10 minutes to answer
AI_TIMEOUT_SECONDS = 600
GAINERS_HEADLINE = "Top gainers"
LOSERS_HEADLINE = "Top losers"
TOPICS = ["Earnings / results", "Product / orders", "Guidance / outlook", "Analyst rating", "Macro / sector",
          "Legal / regulatory", "Other"]

# How each stock appears in the AI's answer, so the Dashboard can pick it apart:
#   - Uber Technologies Inc (UBER:NYSE) +20.00%: [Product / orders] A summary of its latest news.
STOCK_LINE = re.compile(r"^\s*[-•*]\s*.*?\(([A-Za-z0-9.\-]+)(?::[A-Za-z]+)?\)[^:]*:\s*(?:\[([^\]]+)\]\s*)?(.+?)\s*$")

# Latest news: the top Google News result for each top mover, from its RSS feed (no key needed). It is the
# stock's latest news, which may not be why the stock moved. Today's news is tried first, then any.
NEWS_RSS_URL = "https://news.google.com/rss/search?q={}&hl=en-US&gl=US&ceid=US:en"
NEWS_QUERIES = ["{} stock when:1d", "{} stock"]
NEWS_TIMEOUT_SECONDS = 15
# Check Google News's certificate with certifi's certificates when they're installed (they come with requests),
# since some Pythons (like the python.org ones on macOS) can't find any certificates by themselves
try:
    import certifi
    NEWS_SSL_CONTEXT = ssl.create_default_context(cafile=certifi.where())
except ImportError:
    NEWS_SSL_CONTEXT = ssl.create_default_context()

# xlsx layout
DASHBOARD_SHEET_TITLE = "Dashboard"
DATA_SHEET_TITLE = "Stock Data"
REPORT_TITLE = "End-of-day watchlist report"
XLSX_HEADER = ["Stock Name", "Ticker", "Market", "Yesterday's Price", "Today's Price", "Total Difference", "Percent Difference"]
DATA_COLUMN_WIDTHS = [38, 10, 10, 18, 16, 18, 20]
DASHBOARD_HEADER = ["Side", "Ticker", "Move", "Topic", "Latest news", "Source"]
DASHBOARD_COLUMN_WIDTHS = [12, 9, 12, 22, 82, 44]
DASHBOARD_HEADER_ROW = 7
HEADER_NOTES = {
    "Stock Name": "Company name.",
    "Ticker": "The symbol the stock trades under, e.g. AAPL.",
    "Market": "The exchange the stock is listed on (NYSE or Nasdaq). It says little about what moves the stock.",
    "Yesterday's Price": "Usually called the previous close. Use a split-adjusted value, or a stock split will look like a crash.",
    "Today's Price": "Today's closing price (the close), set at the 4:00pm ET bell.",
    "Total Difference": "Change in dollars: close minus previous close.",
    "Percent Difference": "Change ÷ previous close, usually called % change. The ▲/▼ come from the number format, "
                          "so direction reads without relying on red vs. green.",
}
DATA_FOOTNOTE = "Hover any header for a definition."
NO_NEWS = "No recent news found."
NO_SUMMARY = "No summary available."
PLACEHOLDER = "—"

# Number formats: dollars, signed dollars, and a ▲/▼ percent (no change shows a plain 0)
PRICE_FORMAT = "\\$#,##0.00"
DIFFERENCE_FORMAT = "\\+\\$#,##0.00;\\-\\$#,##0.00;\\$0.00"
PERCENT_FORMAT = "\"▲ \"0.00%;\"▼ \"0.00%;0"

# Fonts, borders and alignment
TITLE_FONT = Font(name="Calibri", bold=True, size=16)
NOTE_FONT = Font(name="Calibri", italic=True, size=10, color="52514E")
SUMMARY_FONT = Font(name="Calibri", bold=True, size=14)
BOLD_FONT = Font(name="Calibri", bold=True, size=11)
MUTED_FONT = Font(name="Calibri", italic=True, size=11, color="52514E")
HEADER_BORDER = Border(bottom=Side(style="thin", color="C3C2B7"))
TOP = Alignment(vertical="top")
TOP_WRAP = Alignment(vertical="top", wrap_text=True)

# Percent difference shading: a move near 0 gets the light color, the biggest move of the day gets the dark one
GAIN_LIGHT, GAIN_DARK = (0xE6, 0xF4, 0xEA), (0x1E, 0x7B, 0x34)
LOSS_LIGHT, LOSS_DARK = (0xFC, 0xE8, 0xE6), (0xB3, 0x26, 0x1E)
BLACK, WHITE = "000000", "FFFFFF"

# Emojis (there is no bull face emoji, so the gainers get a cow face)
COW_FACE, BEAR_FACE = "\U0001F42E", "\U0001F43B"
SUN, CLOUD, RAIN = "☀️", "⛅", "\U0001F327️"

logger = logging.getLogger()


def run_analyze(mariadb_data, mariadb_cursor, config):
    '''
    helps to facilitate the running of the analyze functions.

    Parameters:
    -----------
      mariadb_data: the stock data provided from the distributor
      mariadb_cursor: a cursor to mariadb, so the analyzer functions can also use the mariadb reference
      config: the distributor config in use

    Returns:
    -----------
      - email contents that could be returned to the recipient describing the .xlsx file: the Dashboard
        as HTML, from develop_email (None if the data could not be analyzed, so the distributor sends an apology instead)

    Creates:
    -----------
      - a file at the location of email_config -> attachment
    '''
    try:
        # call format_findings
        formatted_data = format_findings(mariadb_data)
        if not formatted_data:
            logger.info("No stock had both a previous and a current price, so there is nothing to analyze")
            return None

        # add each stock's name and acronym from mariadb
        for finding in formatted_data:
            mariadb_cursor.execute(SELECT_STOCK_NAME_AND_ACRONYM.format(finding["stock_id"]))
            rows = mariadb_cursor.fetchall()
            if rows:
                finding["stock_name"], finding["acronym"] = rows[0]

        # use sort_findings
        sorted_data = sort_findings(formatted_data)

        # add each top gainer's and loser's latest news
        gainers, losers = top_movers(sorted_data)
        for finding in gainers + losers:
            finding["news"] = fetch_latest_news(finding)

        # use get_insight
        blurb = get_insight(sorted_data)

        # develop_xlsx
        develop_xlsx(sorted_data, blurb, config[EMAIL_CONFIG][ATTACHMENT])

        # the Dashboard again, as the email body
        return develop_email(sorted_data, blurb)
    except Exception as e:
        logger.error("Analysis failed: {}".format(e))
        traceback.print_exc()
        return None


def format_findings(mariadb_data):
    '''
    Takes in the data from mariadb, and formats it for use within the analyzer functions.

    Parameters:
    -----------
      mariadb_data: the stock data provided from the distributor
        (rows of (stock_id, price), newest pull first)

    Returns:
    -----------
      - a formatted set of findings: a list of dicts, one per stock, with stock_id, stock_name, acronym,
        yesterday_price, today_price, difference and rate_of_change. Stocks without two prices,
        with a previous price of 0, or with either price of -1 (a missing price) are left out,
        so they appear in neither the xlsx nor the blurb.
    '''
    prices_by_stock = {}
    for (stock_id, price) in mariadb_data:
        prices_by_stock.setdefault(stock_id, []).append(price)

    findings = []
    for stock_id, prices in prices_by_stock.items():
        if len(prices) < 2:
            continue
        today_price, yesterday_price = prices[0], prices[1]
        if MISSING_PRICE in (today_price, yesterday_price) or yesterday_price == 0:
            continue
        difference = today_price - yesterday_price
        findings.append({
            "stock_id": stock_id,
            "stock_name": None,
            "acronym": None,
            "yesterday_price": yesterday_price,
            "today_price": today_price,
            "difference": difference,
            "rate_of_change": difference / yesterday_price,
        })
    return findings


def sort_findings(formatted_mariadb_data):
    '''
    Sorts the findings following getting them into a usable format.

    Parameters:
    -----------
      formatted_mariadb_data: the stock data provided from the distributor, formatted

    Returns:
    -----------
      - a formatted set of findings sorted by rate of change (biggest gainer first, biggest loser last)
    '''
    return sorted(formatted_mariadb_data, key=lambda finding: finding["rate_of_change"], reverse=True)


def top_movers(sorted_data):
    '''
    Picks out the top gainers and losers.

    Parameters:
    -----------
      sorted_data: a sorted set of stock data

    Returns:
    -----------
      - (gainers, losers): up to three stocks that rose, biggest gain first, and up to three that fell,
        biggest drop first. Unchanged stocks are in neither.
    '''
    gainers = [finding for finding in sorted_data if finding["rate_of_change"] > 0][:TOP_COUNT]
    losers = [finding for finding in reversed(sorted_data) if finding["rate_of_change"] < 0][:TOP_COUNT]
    return gainers, losers


def fetch_latest_news(finding):
    '''
    Finds a stock's latest news: the top Google News result for it (today's if there is any).

    Parameters:
    -----------
      finding: a stock's finding, with its acronym (or name)

    Returns:
    -----------
      - {"title", "url", "source", "published"} for the top article, or None if nothing was found
        (or Google News couldn't be reached). The url is a Google News link that opens the article.
    '''
    name = ticker_and_market(finding)[0] or finding["stock_name"]
    if not name:
        return None
    for query in NEWS_QUERIES:
        url = NEWS_RSS_URL.format(quote_plus(query.format(name)))
        try:
            request = Request(url, headers={"User-Agent": "Mozilla/5.0 (StockRate Distributor)"})
            with urlopen(request, timeout=NEWS_TIMEOUT_SECONDS, context=NEWS_SSL_CONTEXT) as response:
                item = ElementTree.fromstring(response.read()).find("./channel/item")
        except Exception as e:
            logger.error("Couldn't get the latest news for {}: {}".format(name, e))
            return None
        if item is None:
            continue
        title, link = (item.findtext("title") or "").strip(), (item.findtext("link") or "").strip()
        source = (item.findtext("source") or "").strip()
        # Google News titles end with " - Source"
        if source and title.endswith(" - " + source):
            title = title[:-len(" - " + source)]
        if title and link:
            return {"title": title, "url": link, "source": source or None,
                    "published": (item.findtext("pubDate") or "").strip() or None}
    return None


def get_insight(sorted_data):
    '''
    Talks to an AI model about the current sorted data, specifically asking for insight on the top three gainers
    and losers: a short summary of each one's latest news (from fetch_latest_news, set by run_analyze as
    finding["news"]), with a topic for it.

    Parameters:
    -----------
      sorted_data: a sorted set of stock data

    Returns:
    -----------
      - a blurb for both the email and the xlsx file (AI_ERROR_ANSWER if every model in GEMINI_MODELS fails)
    '''
    gainers, losers = top_movers(sorted_data)

    def describe(findings):
        lines = []
        for finding in findings:
            lines.append("{} ({}): {} -> {} ({:+.2%})".format(finding["stock_name"], finding["acronym"],
                finding["yesterday_price"], finding["today_price"], finding["rate_of_change"]))
            news = finding.get("news")
            if news:
                lines.append("  Latest news: \"{}\" ({}{})".format(news["title"], news["source"] or "unknown source",
                    ", " + news["published"] if news["published"] else ""))
            else:
                lines.append("  Latest news: none found")
        return "\n".join(lines) or "None today"

    query = ("Here are today's top three gainers and top three losers among the stocks we track, each with the\n"
        "headline of its latest news from Google News.\n"
        "For each stock, write one or two sentences summarizing what its latest news is about, based only on the\n"
        "headline given; don't add details that aren't in it. The news may not be why the stock moved, so don't\n"
        "say that it is. If a stock has no news, write \"{}\" for it. Keep the whole answer short enough for an email.\n"
        "Start the gainers section with a line saying only \"{}\" and the losers section with a line saying only \"{}\".\n"
        "Write each stock as exactly one line, in plain text:\n"
        "- Stock Name (TICKER:MARKET) +1.23%: [Topic] One or two sentences about its latest news.\n"
        "where Topic is the news's topic, one of: {}.\n\n"
        "Top gainers:\n{}\n\nTop losers:\n{}\n").format(NO_NEWS, GAINERS_HEADLINE, LOSERS_HEADLINE, ", ".join(TOPICS),
            describe(gainers), describe(losers))

    for attempt, model_name in enumerate(GEMINI_MODELS, start=1):
        try:
            model = genai.GenerativeModel(model_name)
            return model.generate_content(query, request_options={"timeout": AI_TIMEOUT_SECONDS}).text
        except Exception as e:
            logger.error("AI query to {} failed (attempt {} of {}): {}".format(model_name, attempt, len(GEMINI_MODELS), e))
            traceback.print_exc()
            if attempt < len(GEMINI_MODELS):
                wait = AI_RETRY_WAITS[attempt - 1]
                logger.info("Waiting {} seconds before asking {}".format(wait, GEMINI_MODELS[attempt]))
                sleep(wait)
    return AI_ERROR_ANSWER


def parse_insight(blurb):
    '''
    Picks each stock's details out of the AI's blurb, for the Dashboard.

    Parameters:
    -----------
      blurb: the blurb from get_insight

    Returns:
    -----------
      - a dict of ticker (upper case) -> {"topic", "summary"} (the summary of its latest news); a topic the
        AI left out is None, and stocks the AI didn't write about are missing
    '''
    details = {}
    for line in blurb.splitlines():
        stock = STOCK_LINE.match(line.replace("**", "").replace("`", ""))
        if stock:
            details[stock.group(1).upper()] = {"topic": stock.group(2), "summary": stock.group(3)}
    return details


def shade_for_rate(rate_of_change, max_abs_rate):
    '''
    Works out the shading for a percent difference cell.

    Parameters:
    -----------
      rate_of_change: the stock's rate of change
      max_abs_rate: the biggest rate of change (up or down) among all the stocks, which gets the darkest shade

    Returns:
    -----------
      - (fill color, font color) as hex strings: green for a gain and red for a loss, lighter the closer
        the rate is to 0 (with white text on the darker shades), or None for no change
    '''
    if rate_of_change == 0 or max_abs_rate == 0:
        return None
    light, dark = (GAIN_LIGHT, GAIN_DARK) if rate_of_change > 0 else (LOSS_LIGHT, LOSS_DARK)
    strength = min(abs(rate_of_change) / max_abs_rate, 1.0)
    fill = "".join("{:02X}".format(round(l + (d - l) * strength)) for l, d in zip(light, dark))
    return fill, WHITE if strength > 0.5 else BLACK


def describe_move(rate_of_change):
    '''
    A rate of change as text, like the percent number format: "▲ 20.00%", "▼ 25.00%" or "0.00%".
    '''
    arrow = "▲ " if rate_of_change > 0 else "▼ " if rate_of_change < 0 else ""
    return "{}{:.2%}".format(arrow, abs(rate_of_change))


def ticker_and_market(finding):
    '''
    Splits a stock's acronym: "GOOGL:NASDAQ" -> ("GOOGL", "NASDAQ").
    '''
    ticker, _, market = (finding["acronym"] or "").partition(":")
    return ticker, market


def summarize_day(sorted_data):
    '''
    Sums up the day for the top of the Dashboard and the email.

    Parameters:
    -----------
      sorted_data: a sorted set of stock data

    Returns:
    -----------
      - (mood line, biggest move line), like
        ("⛅ Mixed day — 10 up · 9 down · 1 flat · watchlist average ▼ 0.46%",
         "Biggest gain UBER ▲ 20.00% · biggest drop AAPL ▼ 25.00%")
    '''
    rates = [finding["rate_of_change"] for finding in sorted_data]
    up, down = sum(rate > 0 for rate in rates), sum(rate < 0 for rate in rates)
    average = sum(rates) / len(rates) if rates else 0
    if up >= 2 * down and up > 0:
        mood = "{} Up day".format(SUN)
    elif down >= 2 * up and down > 0:
        mood = "{} Down day".format(RAIN)
    else:
        mood = "{} Mixed day".format(CLOUD)
    mood_line = "{} — {} up · {} down · {} flat · watchlist average {}".format(
        mood, up, down, len(rates) - up - down, describe_move(average))

    gainers, losers = top_movers(sorted_data)
    biggest_gain = "{} {}".format(ticker_and_market(gainers[0])[0], describe_move(gainers[0]["rate_of_change"])) \
        if gainers else PLACEHOLDER
    biggest_drop = "{} {}".format(ticker_and_market(losers[0])[0], describe_move(losers[0]["rate_of_change"])) \
        if losers else PLACEHOLDER
    return mood_line, "Biggest gain {} · biggest drop {}".format(biggest_gain, biggest_drop)


def develop_email(sorted_data, blurb):
    '''
    Develops the email body: the Dashboard as simple HTML that reads well on a phone.

    Parameters:
    -----------
      sorted_data: a sorted set of stock data
      blurb: the blurb from get_insight

    Returns:
    -----------
      - the email body as one line of HTML (yagmail turns line breaks into <br>, so there are none): the title
        and date, the summary of the day, the biggest gain and drop, then each top gainer and loser with its
        move, the topic and summary of its latest news, and a link to that news
    '''
    esc = html.escape
    mood_line, biggest_line = summarize_day(sorted_data)
    gainers, losers = top_movers(sorted_data)
    details = parse_insight(blurb)

    parts = ['<div style="font-family:-apple-system,Helvetica,Arial,sans-serif;max-width:600px;color:#222;">',
             '<div style="font-size:20px;font-weight:bold;">{}</div>'.format(esc(REPORT_TITLE)),
             '<div style="font-size:12px;font-style:italic;color:#52514E;">Report for {}</div>'.format(
                 datetime.date.today().isoformat()),
             '<div style="font-size:16px;font-weight:bold;margin-top:12px;">{}</div>'.format(esc(mood_line)),
             '<div style="font-size:14px;margin-top:4px;">{}</div>'.format(esc(biggest_line))]
    if blurb == AI_ERROR_ANSWER:
        parts.append('<div style="font-size:13px;font-style:italic;color:#52514E;margin-top:8px;">{}</div>'.format(
            esc(AI_ERROR_ANSWER)))

    for (emoji, headline, findings) in [(COW_FACE, GAINERS_HEADLINE, gainers), (BEAR_FACE, LOSERS_HEADLINE, losers)]:
        if not findings:
            continue
        parts.append('<div style="font-size:16px;font-weight:bold;margin-top:18px;border-bottom:1px solid #C3C2B7;'
                     'padding-bottom:4px;">{} {}</div>'.format(emoji, esc(headline)))
        for finding in findings:
            ticker = ticker_and_market(finding)[0]
            detail = details.get(ticker.upper(), {})
            color = "#1E7B34" if finding["rate_of_change"] > 0 else "#B3261E"
            parts.append('<div style="margin-top:10px;">')
            parts.append('<div style="font-size:15px;"><b>{}</b> <span style="color:{};font-weight:bold;">{}</span>'
                         '<span style="color:#52514E;"> · {}</span></div>'.format(
                             esc(ticker), color, esc(describe_move(finding["rate_of_change"])),
                             esc(detail.get("topic") or PLACEHOLDER)))
            news = finding.get("news")
            muted = '<div style="font-size:14px;margin-top:2px;font-style:italic;color:#52514E;">{}</div>'
            if not news:
                parts.append(muted.format(esc(NO_NEWS)))
            else:
                if detail.get("summary"):
                    parts.append('<div style="font-size:14px;margin-top:2px;"><b>Latest {} news:</b> {}</div>'.format(
                        esc(ticker), esc(detail["summary"])))
                else:
                    parts.append(muted.format(esc(NO_SUMMARY)))
                parts.append('<div style="font-size:13px;margin-top:2px;"><a href="{}">{}</a>{}</div>'.format(
                    esc(news["url"], quote=True), esc(news["title"]),
                    '<span style="color:#52514E;"> · {}</span>'.format(esc(news["source"])) if news["source"] else ""))
            parts.append('</div>')

    parts.append('<div style="font-size:12px;color:#52514E;margin-top:18px;">All {} stocks are in the attached '
                 'spreadsheet.</div>'.format(len(sorted_data)))
    parts.append('</div>')
    return "".join(parts)


def develop_xlsx(sorted_data, blurb, xlsx_location):
    '''
    Uses information gathered earlier in the function calls and develops an xlsx from it.

    Parameters:
    -----------
      sorted_data: a sorted set of stock data
      blurb: a blurb to be included in the xlsx
      xlsx_location: the location of the xlsx file

    Creates:
    -----------
      - a file at the location of email_config -> attachment, with two sheets:
        "Dashboard": the title and date, a summary of the day, the biggest gain and drop, and a table of the
          top gainers and losers with the topic and summary of its latest news (from the blurb) and a link to it
        "Stock Data": one row per stock in alphabetical order by market then name, with its acronym split
          into ticker and market, prices in dollars, and its percent difference marked ▲/▼ and shaded green
          or red (darker the bigger the move); the header row is frozen, filterable and has hover notes
    '''
    workbook = openpyxl.Workbook()
    dashboard = workbook.active
    dashboard.title = DASHBOARD_SHEET_TITLE
    data_sheet = workbook.create_sheet(DATA_SHEET_TITLE)
    for sheet in (dashboard, data_sheet):
        sheet.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
        sheet.page_setup.fitToHeight = 0
        sheet.page_setup.orientation = "landscape"

    # ---- Dashboard: title, date and summary of the day
    for column, width in zip("ABCDEF", DASHBOARD_COLUMN_WIDTHS):
        dashboard.column_dimensions[column].width = width
    dashboard["A1"] = REPORT_TITLE
    dashboard["A1"].font = TITLE_FONT
    dashboard["A2"] = "Report for {}".format(datetime.date.today().isoformat())
    dashboard["A2"].font = NOTE_FONT

    mood_line, biggest_line = summarize_day(sorted_data)
    dashboard["A3"] = mood_line
    dashboard["A3"].font = SUMMARY_FONT
    dashboard["A4"] = biggest_line
    gainers, losers = top_movers(sorted_data)

    # ---- Dashboard: the top gainers and losers, with details from the blurb
    for column, header in enumerate(DASHBOARD_HEADER, start=1):
        cell = dashboard.cell(row=DASHBOARD_HEADER_ROW, column=column, value=header)
        cell.font = BOLD_FONT
        cell.border = HEADER_BORDER

    details = parse_insight(blurb)
    movers = [("{} Gainer".format(COW_FACE), finding) for finding in gainers] + \
             [("{} Loser".format(BEAR_FACE), finding) for finding in losers]
    for row, (side, finding) in enumerate(movers, start=DASHBOARD_HEADER_ROW + 1):
        ticker = ticker_and_market(finding)[0]
        detail = details.get(ticker.upper(), {})
        side_cell = dashboard.cell(row=row, column=1, value=side)
        side_cell.font = BOLD_FONT
        side_cell.alignment = TOP
        dashboard.cell(row=row, column=2, value=ticker).alignment = TOP
        move_cell = dashboard.cell(row=row, column=3, value=finding["rate_of_change"])
        move_cell.number_format = PERCENT_FORMAT
        move_cell.alignment = TOP
        dashboard.cell(row=row, column=4, value=detail.get("topic") or PLACEHOLDER).alignment = TOP

        news = finding.get("news")
        summary = detail.get("summary") if news else None
        news_cell = dashboard.cell(row=row, column=5, value=summary or (NO_SUMMARY if news else NO_NEWS))
        news_cell.alignment = TOP_WRAP
        if not summary:
            news_cell.font = MUTED_FONT

        source_cell = dashboard.cell(row=row, column=6, value=PLACEHOLDER)
        if news:
            source_cell.value = "{} ({})".format(news["title"], news["source"]) if news["source"] else news["title"]
            source_cell.hyperlink = news["url"]
            source_cell.style = "Hyperlink"
        source_cell.alignment = TOP_WRAP

    # ---- Stock Data: every stock, alphabetical by market then name
    data_sheet.append(XLSX_HEADER)
    for cell in data_sheet[1]:
        cell.font = BOLD_FONT
        cell.comment = Comment(HEADER_NOTES[cell.value], "StockRate")
    for column, width in zip("ABCDEFG", DATA_COLUMN_WIDTHS):
        data_sheet.column_dimensions[column].width = width

    rows = sorted(sorted_data, key=lambda finding: (ticker_and_market(finding)[1].casefold(),
                                                    (finding["stock_name"] or "").casefold()))
    max_abs_rate = max((abs(finding["rate_of_change"]) for finding in sorted_data), default=0)
    for finding in rows:
        ticker, market = ticker_and_market(finding)
        data_sheet.append([finding["stock_name"] or "", ticker, market, finding["yesterday_price"],
            finding["today_price"], finding["difference"], finding["rate_of_change"]])
        row = data_sheet.max_row
        data_sheet.cell(row=row, column=4).number_format = PRICE_FORMAT
        data_sheet.cell(row=row, column=5).number_format = PRICE_FORMAT
        data_sheet.cell(row=row, column=6).number_format = DIFFERENCE_FORMAT
        percent_cell = data_sheet.cell(row=row, column=7)
        percent_cell.number_format = PERCENT_FORMAT
        shade = shade_for_rate(finding["rate_of_change"], max_abs_rate)
        if shade is None:
            # no change: just a black 0
            percent_cell.font = Font(name="Calibri", size=11, color=BLACK)
        else:
            fill, font = shade
            percent_cell.fill = PatternFill(fill_type="solid", start_color=fill, end_color=fill)
            percent_cell.font = Font(name="Calibri", size=11, color=font)

    last_row = data_sheet.max_row
    data_sheet.freeze_panes = "B2"
    data_sheet.auto_filter.ref = "A1:G{}".format(last_row)
    data_sheet.cell(row=last_row + 2, column=1, value=DATA_FOOTNOTE).font = NOTE_FONT

    workbook.save(xlsx_location)
