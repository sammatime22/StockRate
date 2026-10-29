# Unit tests for analyzer.py, one per function. No MariaDB or Gemini is needed.
# sammatime22, 2026
#
# Setup on iSH (Alpine):
#   apk add python3 py3-pytest py3-openpyxl
#
# Run from the Distributor directory:
#   python3 -m pytest test_analyzer.py
import datetime
import html
import io
import os
import sys
import types
from urllib.parse import unquote_plus

import openpyxl
import pytest

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, HERE)

# analyzer imports google.generativeai at the top, so stand it in before importing analyzer
if "google.generativeai" not in sys.modules:
    fake_genai = types.ModuleType("google.generativeai")
    fake_genai.GenerativeModel = None
    fake_google = types.ModuleType("google")
    fake_google.generativeai = fake_genai
    sys.modules["google"] = fake_google
    sys.modules["google.generativeai"] = fake_genai

import analyzer

# (stock_id, stock_name, acronym, yesterday's price, today's price); every rate of change is different
STOCKS = [
    (1, "Alphabet Inc Class A", "GOOGL:NASDAQ", 100.0, 110.0),      # +10%
    (2, "Apple Inc", "AAPL:NASDAQ", 200.0, 150.0),                  # -25%
    (3, "Microsoft Corp", "MSFT:NASDAQ", 50.0, 50.0),               # 0%
    (4, "Amazon.com Inc", "AMZN:NASDAQ", 180.0, 189.0),             # +5%
    (5, "NVIDIA Corp", "NVDA:NASDAQ", 120.0, 138.0),                # +15%
    (6, "Meta Platforms Inc", "META:NASDAQ", 500.0, 490.0),         # -2%
    (7, "Tesla Inc", "TSLA:NASDAQ", 250.0, 200.0),                  # -20%
    (8, "Netflix Inc", "NFLX:NASDAQ", 600.0, 612.0),                # +2%
    (9, "Advanced Micro Devices Inc", "AMD:NASDAQ", 160.0, 164.0),  # +2.5%
    (10, "Intel Corp", "INTC:NASDAQ", 30.0, 27.0),                  # -10%
    (11, "International Business Machines Corp", "IBM:NYSE", 200.0, 207.0),  # +3.5%
    (12, "Oracle Corp", "ORCL:NYSE", 140.0, 133.0),                 # -5%
    (13, "Salesforce Inc", "CRM:NYSE", 300.0, 324.0),               # +8%
    (14, "Adobe Inc", "ADBE:NASDAQ", 500.0, 470.0),                 # -6%
    (15, "PayPal Holdings Inc", "PYPL:NASDAQ", 80.0, 79.0),         # -1.25%
    (16, "Uber Technologies Inc", "UBER:NYSE", 70.0, 84.0),         # +20%
    (17, "Walt Disney Co", "DIS:NYSE", 100.0, 97.0),                # -3%
    (18, "Coca-Cola Co", "KO:NYSE", 200.0, 202.0),                  # +1%
    (19, "JPMorgan Chase & Co", "JPM:NYSE", 150.0, 138.0),          # -8%
    (20, "Walmart Inc", "WMT:NYSE", 100.0, 104.0),                  # +4%
]
# Canned latest news for the top movers, as Google News would give it: ticker -> (headline, source, link).
# The headlines and links are made up (news.example.com), not real news.
CANNED_NEWS = {
    "UBER": ("Uber Takes Robotaxis to Ten More Cities", "Example News", "https://news.example.com/2026/10/uber-robotaxi-expansion"),
    "NVDA": ("Cloud Giants Line Up for NVIDIA's Next Chip", "Example Wire", "https://news.example.com/2026/10/nvidia-data-center-orders"),
    "GOOGL": ("Alphabet Earnings Beat on Cloud Growth", "Example News", "https://news.example.com/2026/10/alphabet-earnings-beat"),
    "AAPL": ("Apple Delays Launch Amid Soft Demand", "Example Wire", "https://news.example.com/2026/10/apple-launch-delay"),
    "TSLA": ("Tesla Deliveries Miss Estimates", "Example News", "https://news.example.com/2026/10/tesla-deliveries-miss"),
    "INTC": ("Intel Lowers Guidance on PC Slump", "Example Wire", "https://news.example.com/2026/10/intel-guidance-cut"),
}
NEWS_DATE = "Wed, 28 Oct 2026 14:00:00 GMT"

# A canned Gemini answer for the 20 test stocks: a summary of each top mover's (made-up) latest news
CANNED_INSIGHT = """Today's biggest movers (canned test insight - summaries of made-up news)

Top gainers
- Uber Technologies Inc (UBER:NYSE) +20.00%: [Product / orders] Uber is expanding its robotaxi service to ten more cities.
- NVIDIA Corp (NVDA:NASDAQ) +15.00%: [Product / orders] Several cloud providers are lining up to buy NVIDIA's next chip.
- Alphabet Inc Class A (GOOGL:NASDAQ) +10.00%: [Earnings / results] Alphabet's earnings beat expectations on cloud growth.

Top losers
- Apple Inc (AAPL:NASDAQ) -25.00%: [Product / orders] Apple delayed a product launch amid soft demand.
- Tesla Inc (TSLA:NASDAQ) -20.00%: [Earnings / results] Tesla's deliveries missed estimates.
- Intel Corp (INTC:NASDAQ) -10.00%: [Guidance / outlook] Intel lowered its guidance because of a PC slump.
"""

STOCK_NAMES = {stock_id: (name, acronym) for (stock_id, name, acronym, _, _) in STOCKS}

# Biggest gainer first, biggest loser last
EXPECTED_ORDER = ["UBER:NYSE", "NVDA:NASDAQ", "GOOGL:NASDAQ", "CRM:NYSE", "AMZN:NASDAQ", "WMT:NYSE", "IBM:NYSE",
                  "AMD:NASDAQ", "NFLX:NASDAQ", "KO:NYSE", "MSFT:NASDAQ", "PYPL:NASDAQ", "META:NASDAQ", "DIS:NYSE",
                  "ORCL:NYSE", "ADBE:NASDAQ", "JPM:NYSE", "INTC:NASDAQ", "TSLA:NASDAQ", "AAPL:NASDAQ"]

# The xlsx lists stocks alphabetically by market, then by name
XLSX_ORDER = ["ADBE:NASDAQ", "AMD:NASDAQ", "GOOGL:NASDAQ", "AMZN:NASDAQ", "AAPL:NASDAQ", "INTC:NASDAQ", "META:NASDAQ",
              "MSFT:NASDAQ", "NFLX:NASDAQ", "NVDA:NASDAQ", "PYPL:NASDAQ", "TSLA:NASDAQ",
              "KO:NYSE", "IBM:NYSE", "JPM:NYSE", "ORCL:NYSE", "CRM:NYSE", "UBER:NYSE", "WMT:NYSE", "DIS:NYSE"]

COW_FACE, BEAR_FACE = "\U0001F42E", "\U0001F43B"

# The Dashboard's gainers/losers table for CANNED_NEWS and CANNED_INSIGHT:
# (side, ticker, move, topic, latest news summary, source, link)
EXPECTED_DASHBOARD = [
    (side, ticker, move, topic, summary, "{} ({})".format(CANNED_NEWS[ticker][0], CANNED_NEWS[ticker][1]),
     CANNED_NEWS[ticker][2])
    for (side, ticker, move, topic, summary) in [
        (COW_FACE + " Gainer", "UBER", 0.2, "Product / orders", "Uber is expanding its robotaxi service to ten more cities."),
        (COW_FACE + " Gainer", "NVDA", 0.15, "Product / orders", "Several cloud providers are lining up to buy NVIDIA's next chip."),
        (COW_FACE + " Gainer", "GOOGL", 0.1, "Earnings / results", "Alphabet's earnings beat expectations on cloud growth."),
        (BEAR_FACE + " Loser", "AAPL", -0.25, "Product / orders", "Apple delayed a product launch amid soft demand."),
        (BEAR_FACE + " Loser", "TSLA", -0.2, "Earnings / results", "Tesla's deliveries missed estimates."),
        (BEAR_FACE + " Loser", "INTC", -0.1, "Guidance / outlook", "Intel lowered its guidance because of a PC slump."),
    ]
]

# Newest pull first, as the distributor's query returns it: today's pull then yesterday's
MARIADB_DATA = ([(stock_id, today) for (stock_id, _, _, _, today) in reversed(STOCKS)] +
                [(stock_id, yesterday) for (stock_id, _, _, yesterday, _) in reversed(STOCKS)])


class FakeCursor:
    '''
    Answers the analyzer's stock name query from STOCK_NAMES.
    '''

    def __init__(self):
        self.stock_id = None

    def execute(self, sql):
        self.stock_id = int(sql.rsplit("=", 1)[1].rstrip(";"))

    def fetchall(self):
        return [STOCK_NAMES[self.stock_id]] if self.stock_id in STOCK_NAMES else []


class FakeGenAI:
    '''
    Records the question asked (and its request options) and gives a canned answer;
    the first `failures` questions raise instead, like a 504 Deadline Exceeded.
    '''

    def __init__(self, answer, failures=0):
        self.answer = answer
        self.failures = failures
        self.questions = []
        self.request_options = []
        self.models = []

    def GenerativeModel(self, model_name):
        fake = self
        fake.models.append(model_name)

        class Model:
            def generate_content(self, query, request_options=None):
                fake.questions.append(query)
                fake.request_options.append(request_options)
                if len(fake.questions) <= fake.failures:
                    raise RuntimeError("504 Deadline Exceeded")
                return types.SimpleNamespace(text=fake.answer)

        return Model()


def rss(items):
    '''
    A Google News RSS feed with these (headline, source, link) items; like Google News, each title ends in " - Source".
    '''
    entries = "".join("<item><title>{} - {}</title><link>{}</link><pubDate>{}</pubDate>"
                      "<source url=\"https://example.com\">{}</source></item>".format(
                          html.escape(title), html.escape(source), html.escape(link), NEWS_DATE, html.escape(source))
                      for (title, source, link) in items)
    return '<?xml version="1.0"?><rss version="2.0"><channel><title>Google News</title>{}</channel></rss>'.format(entries)


class FakeGoogleNews:
    '''
    Stands in for urlopen: answers Google News RSS searches from `today` (searches for today's news) and
    `any_time` (searches for any news), both ticker -> list of (headline, source, link), recording each request;
    fail=True makes every request raise.
    '''

    def __init__(self, today=None, any_time=None, fail=False):
        self.today = today if today is not None else {ticker: [news] for ticker, news in CANNED_NEWS.items()}
        self.any_time = any_time or {}
        self.fail = fail
        self.requests = []
        self.contexts = []

    def __call__(self, request, timeout=None, context=None):
        self.requests.append(request)
        self.contexts.append(context)
        if self.fail:
            raise OSError("Google News made to fail by the test")
        query = unquote_plus(request.full_url.split("q=")[1].split("&")[0])
        ticker = query.split()[0]
        items = self.today.get(ticker, []) if "when:1d" in query else self.any_time.get(ticker, [])
        return io.BytesIO(rss(items).encode())

    def queries(self):
        return [unquote_plus(request.full_url.split("q=")[1].split("&")[0]) for request in self.requests]


def news_for(ticker):
    '''
    The latest news fetch_latest_news should find for a ticker in CANNED_NEWS.
    '''
    title, source, url = CANNED_NEWS[ticker]
    return {"title": title, "url": url, "source": source, "published": NEWS_DATE}


def sorted_with_news():
    '''
    All the test stocks, sorted, with each top mover's latest news added (as run_analyze does).
    '''
    sorted_data = analyzer.sort_findings([finding(stock_id) for (stock_id, _, _, _, _) in STOCKS])
    gainers, losers = analyzer.top_movers(sorted_data)
    for f in gainers + losers:
        f["news"] = news_for(analyzer.ticker_and_market(f)[0])
    return sorted_data


def finding(stock_id, with_names=True):
    '''
    The finding format_findings should produce for a stock in STOCKS (with names, as run_analyze adds them).
    '''
    (_, name, acronym, yesterday_price, today_price) = STOCKS[stock_id - 1]
    return {"stock_id": stock_id, "stock_name": name if with_names else None, "acronym": acronym if with_names else None,
            "yesterday_price": yesterday_price, "today_price": today_price, "difference": today_price - yesterday_price,
            "rate_of_change": (today_price - yesterday_price) / yesterday_price}


def xlsx_row(acronym):
    f = next(finding(stock_id) for (stock_id, _, stock_acronym, _, _) in STOCKS if stock_acronym == acronym)
    ticker, market = acronym.split(":")
    return (f["stock_name"], ticker, market, f["yesterday_price"], f["today_price"], f["difference"], f["rate_of_change"])


def test_run_analyze(monkeypatch, tmp_path):
    fake_genai = FakeGenAI(CANNED_INSIGHT)
    monkeypatch.setattr(analyzer, "genai", fake_genai)
    google_news = FakeGoogleNews()
    monkeypatch.setattr(analyzer, "urlopen", google_news)
    xlsx_location = str(tmp_path / "report.xlsx")
    config = {"email_config": {"attachment": xlsx_location}}

    blurb = analyzer.run_analyze(MARIADB_DATA, FakeCursor(), config)

    # the latest news is looked up for just the six top movers, and Gemini is asked about it
    assert [query.split()[0] for query in google_news.queries()] == ["UBER", "NVDA", "GOOGL", "AAPL", "TSLA", "INTC"]
    assert len(fake_genai.questions) == 1 and 'Latest news: "Uber Takes Robotaxis to Ten More Cities"' in fake_genai.questions[0]
    # the email body is the Dashboard as HTML
    assert blurb == analyzer.develop_email(sorted_with_news(), CANNED_INSIGHT)
    workbook = openpyxl.load_workbook(xlsx_location)
    rows = list(workbook["Stock Data"].iter_rows(values_only=True))[:len(STOCKS) + 1]
    assert rows == [tuple(analyzer.XLSX_HEADER)] + [xlsx_row(acronym) for acronym in XLSX_ORDER]
    assert [row[:6] for row in workbook["Dashboard"].iter_rows(min_row=8, max_row=13, values_only=True)] == \
        [expected[:6] for expected in EXPECTED_DASHBOARD]


def test_format_findings():
    # Stock 21 has only one price and stock 22 had a price of 0, so neither can have a rate of change.
    # Stocks 23 and 24 have a missing (-1) old or new price, and stock 25 uses -1 as an int, so all are left out.
    findings = analyzer.format_findings(MARIADB_DATA + [(21, 10.0), (22, 3.0), (22, 0.0), (23, -1.0), (23, 20.0),
                                                        (24, 30.0), (24, -1.0), (25, -1), (25, 40.0)])

    assert len(findings) == len(STOCKS)
    assert {f["stock_id"]: f for f in findings} == {stock_id: finding(stock_id, with_names=False)
                                                    for (stock_id, _, _, _, _) in STOCKS}


def test_sort_findings():
    findings = [finding(stock_id) for (stock_id, _, _, _, _) in STOCKS]

    sorted_findings = analyzer.sort_findings(findings)

    assert [f["acronym"] for f in sorted_findings] == EXPECTED_ORDER


def test_top_movers():
    sorted_data = analyzer.sort_findings([finding(stock_id) for (stock_id, _, _, _, _) in STOCKS])

    gainers, losers = analyzer.top_movers(sorted_data)

    assert [f["acronym"] for f in gainers] == ["UBER:NYSE", "NVDA:NASDAQ", "GOOGL:NASDAQ"]
    assert [f["acronym"] for f in losers] == ["AAPL:NASDAQ", "TSLA:NASDAQ", "INTC:NASDAQ"]
    # an unchanged stock is neither, even when fewer than three stocks moved that way
    assert analyzer.top_movers([finding(1), finding(3)]) == ([finding(1)], [])


def test_fetch_latest_news(monkeypatch):
    google_news = FakeGoogleNews()
    monkeypatch.setattr(analyzer, "urlopen", google_news)

    # today's top Google News result, with " - Source" taken off the headline
    assert analyzer.fetch_latest_news(finding(16)) == news_for("UBER")
    assert google_news.queries() == ["UBER stock when:1d"]
    assert google_news.requests[0].full_url.startswith("https://news.google.com/rss/search?q=UBER+stock+when%3A1d&")
    assert google_news.requests[0].get_header("User-agent")
    assert google_news.contexts == [analyzer.NEWS_SSL_CONTEXT]

    # only the first (top) result is used
    google_news = FakeGoogleNews(today={"UBER": [CANNED_NEWS["UBER"], CANNED_NEWS["NVDA"]]})
    monkeypatch.setattr(analyzer, "urlopen", google_news)
    assert analyzer.fetch_latest_news(finding(16)) == news_for("UBER")

    # with no news today, the latest news from any time is used
    google_news = FakeGoogleNews(today={}, any_time={"UBER": [CANNED_NEWS["UBER"]]})
    monkeypatch.setattr(analyzer, "urlopen", google_news)
    assert analyzer.fetch_latest_news(finding(16)) == news_for("UBER")
    assert google_news.queries() == ["UBER stock when:1d", "UBER stock"]

    # no news at all, or Google News can't be reached: None
    monkeypatch.setattr(analyzer, "urlopen", FakeGoogleNews(today={}))
    assert analyzer.fetch_latest_news(finding(16)) is None
    monkeypatch.setattr(analyzer, "urlopen", FakeGoogleNews(fail=True))
    assert analyzer.fetch_latest_news(finding(16)) is None

    # a stock without an acronym is searched by name
    google_news = FakeGoogleNews(today={})
    monkeypatch.setattr(analyzer, "urlopen", google_news)
    analyzer.fetch_latest_news(dict(finding(1), acronym=None))
    assert google_news.queries()[0] == "Alphabet Inc Class A stock when:1d"


def test_get_insight(monkeypatch):
    fake_genai = FakeGenAI(CANNED_INSIGHT)
    monkeypatch.setattr(analyzer, "genai", fake_genai)
    sorted_data = sorted_with_news()

    blurb = analyzer.get_insight(sorted_data)

    assert blurb == CANNED_INSIGHT
    question = fake_genai.questions[0]
    assert 'a line saying only "Top gainers"' in question and 'a line saying only "Top losers"' in question
    assert "[Topic]" in question and "Earnings / results" in question
    # Gemini summarizes each stock's latest news, without saying it's why the stock moved
    assert "based only on the" in question and "may not be why the stock moved" in question
    assert '  Latest news: "Uber Takes Robotaxis to Ten More Cities" (Example News, {})'.format(NEWS_DATE) in question
    gainers, losers = question.split("\nTop losers:")
    assert "Uber Technologies Inc (UBER:NYSE): 70.0 -> 84.0 (+20.00%)" in gainers
    assert "NVIDIA Corp (NVDA:NASDAQ): 120.0 -> 138.0 (+15.00%)" in gainers
    assert "Alphabet Inc Class A (GOOGL:NASDAQ): 100.0 -> 110.0 (+10.00%)" in gainers
    assert "Apple Inc (AAPL:NASDAQ): 200.0 -> 150.0 (-25.00%)" in losers
    assert "Tesla Inc (TSLA:NASDAQ): 250.0 -> 200.0 (-20.00%)" in losers
    assert "Intel Corp (INTC:NASDAQ): 30.0 -> 27.0 (-10.00%)" in losers
    # only the top three of each are asked about
    assert "Salesforce" not in question
    assert "JPMorgan" not in question
    assert "Microsoft" not in question
    # Gemini 3.8 Flash is asked first, and gets up to 10 minutes to answer
    assert fake_genai.models == ["gemini-3.8-flash"] and fake_genai.request_options == [{"timeout": 600}]

    # a stock without news is still asked about, saying none was found
    fake_genai = FakeGenAI(CANNED_INSIGHT)
    monkeypatch.setattr(analyzer, "genai", fake_genai)
    analyzer.get_insight(analyzer.sort_findings([finding(stock_id) for (stock_id, _, _, _, _) in STOCKS]))
    assert fake_genai.questions[0].count("  Latest news: none found") == 6

    # if a model fails (like a 503 high demand), the next one is asked after waiting (recorded here, not really waited)...
    waits = []
    monkeypatch.setattr(analyzer, "sleep", waits.append)
    fake_genai = FakeGenAI(CANNED_INSIGHT, failures=3)
    monkeypatch.setattr(analyzer, "genai", fake_genai)
    assert analyzer.get_insight(sorted_data) == CANNED_INSIGHT
    assert fake_genai.models == ["gemini-3.8-flash", "gemini-3.7-flash", "gemini-3.6-flash", "gemini-3.1-flash-lite"]
    assert waits == [30, 60, 90]
    # ...but if all four fail, the error answer is used, with no wait after the last
    waits.clear()
    fake_genai = FakeGenAI(CANNED_INSIGHT, failures=4)
    monkeypatch.setattr(analyzer, "genai", fake_genai)
    assert analyzer.get_insight(sorted_data) == analyzer.AI_ERROR_ANSWER
    assert len(fake_genai.questions) == 4 and waits == [30, 60, 90]


def test_parse_insight():
    details = analyzer.parse_insight(CANNED_INSIGHT)

    assert list(details) == ["UBER", "NVDA", "GOOGL", "AAPL", "TSLA", "INTC"]
    for (_, ticker, _, topic, summary, _, _) in EXPECTED_DASHBOARD:
        assert details[ticker] == {"topic": topic, "summary": summary}

    # markdown is ignored, and a stock without a topic still gets its summary
    assert analyzer.parse_insight("* **Tesla Inc (TSLA:NASDAQ) -20.00%:** Deliveries missed.") == \
        {"TSLA": {"topic": None, "summary": "Deliveries missed."}}
    assert analyzer.parse_insight(analyzer.AI_ERROR_ANSWER) == {}


def test_describe_move():
    assert analyzer.describe_move(0.2) == "▲ 20.00%"
    assert analyzer.describe_move(-0.004625) == "▼ 0.46%"
    assert analyzer.describe_move(0.0) == "0.00%"


def test_ticker_and_market():
    assert analyzer.ticker_and_market(finding(1)) == ("GOOGL", "NASDAQ")
    assert analyzer.ticker_and_market(finding(1, with_names=False)) == ("", "")


def test_summarize_day():
    sorted_data = analyzer.sort_findings([finding(stock_id) for (stock_id, _, _, _, _) in STOCKS])

    assert analyzer.summarize_day(sorted_data) == (
        "\u26c5 Mixed day — 10 up · 9 down · 1 flat · watchlist average ▼ 0.46%",
        "Biggest gain UBER ▲ 20.00% · biggest drop AAPL ▼ 25.00%")
    # mostly up, mostly down, and nothing to report
    assert analyzer.summarize_day([finding(16), finding(1), finding(2)])[0].startswith("\u2600\ufe0f Up day — 2 up · 1 down")
    assert analyzer.summarize_day([finding(2), finding(7), finding(1)])[0].startswith("\U0001F327\ufe0f Down day — 1 up · 2 down")
    assert analyzer.summarize_day([finding(3)]) == ("\u26c5 Mixed day — 0 up · 0 down · 1 flat · watchlist average 0.00%",
                                                    "Biggest gain — · biggest drop —")


def test_develop_email():
    sorted_data = sorted_with_news()

    body = analyzer.develop_email(sorted_data, CANNED_INSIGHT)

    # one line of HTML, since yagmail turns line breaks into <br>
    assert "\n" not in body and body.startswith("<div") and body.endswith("</div>")
    assert "End-of-day watchlist report" in body
    assert "Report for {}".format(datetime.date.today().isoformat()) in body
    assert "\u26c5 Mixed day — 10 up · 9 down · 1 flat · watchlist average ▼ 0.46%" in body
    assert "Biggest gain UBER ▲ 20.00% · biggest drop AAPL ▼ 25.00%" in body
    assert body.index(COW_FACE + " Top gainers") < body.index(BEAR_FACE + " Top losers")
    # each mover in order, with its move, topic, latest news summary, and a link to that news (& and ' are escaped)
    positions = []
    for (_, ticker, move, topic, summary, _, url) in EXPECTED_DASHBOARD:
        positions.append(body.index("<b>{}</b>".format(ticker)))
        assert analyzer.describe_move(move) in body and topic in body
        assert "<b>Latest {} news:</b> {}".format(ticker, html.escape(summary)) in body
        title, source, _ = CANNED_NEWS[ticker]
        assert '<a href="{}">{}</a><span style="color:#52514E;"> · {}</span>'.format(url, html.escape(title), source) in body
    assert positions == sorted(positions)
    assert "Salesforce" not in body and "All 20 stocks are in the attached spreadsheet." in body
    assert analyzer.AI_ERROR_ANSWER not in body

    # if the AI failed, the movers and their news links are still listed, with a note
    body = analyzer.develop_email(sorted_data, analyzer.AI_ERROR_ANSWER)
    assert analyzer.AI_ERROR_ANSWER in body and "<b>UBER</b>" in body
    assert body.count("No summary available.") == 6 and body.count("<a href") == 6

    # a stock without news says so, with no link
    sorted_data = analyzer.sort_findings([finding(stock_id) for (stock_id, _, _, _, _) in STOCKS])
    body = analyzer.develop_email(sorted_data, CANNED_INSIGHT)
    assert body.count("No recent news found.") == 6 and "<a href" not in body and "Latest UBER news" not in body


def test_shade_for_rate():
    # the biggest move gets the darkest shade, with white text
    assert analyzer.shade_for_rate(0.25, 0.25) == ("1E7B34", "FFFFFF")
    assert analyzer.shade_for_rate(-0.25, 0.25) == ("B3261E", "FFFFFF")
    # a small move is close to the light shade, with black text
    assert analyzer.shade_for_rate(0.0025, 0.25) == ("E4F3E8", "000000")
    assert analyzer.shade_for_rate(-0.0025, 0.25) == ("FBE6E4", "000000")
    # halfway is halfway between light and dark
    assert analyzer.shade_for_rate(0.125, 0.25) == ("82B88F", "000000")
    # no change, or nothing moved at all, gets no shading
    assert analyzer.shade_for_rate(0.0, 0.25) is None
    assert analyzer.shade_for_rate(0.0, 0.0) is None


def test_develop_xlsx(tmp_path):
    xlsx_location = str(tmp_path / "report.xlsx")
    sorted_data = sorted_with_news()

    analyzer.develop_xlsx(sorted_data, CANNED_INSIGHT, xlsx_location)

    workbook = openpyxl.load_workbook(xlsx_location)
    assert workbook.sheetnames == ["Dashboard", "Stock Data"]

    # ---- Dashboard: title, date, summary of the day, biggest gain and drop
    dashboard = workbook["Dashboard"]
    assert dashboard["A1"].value == "End-of-day watchlist report" and dashboard["A1"].font.bold
    assert dashboard["A2"].value == "Report for {}".format(datetime.date.today().isoformat())
    assert dashboard["A3"].value == "\u26c5 Mixed day — 10 up · 9 down · 1 flat · watchlist average ▼ 0.46%"
    assert dashboard["A4"].value == "Biggest gain UBER ▲ 20.00% · biggest drop AAPL ▼ 25.00%"

    # ---- Dashboard: the gainers/losers table, with each one's latest news and a link to it
    assert [cell.value for cell in dashboard[7]] == ["Side", "Ticker", "Move", "Topic", "Latest news", "Source"]
    assert all(cell.font.bold for cell in dashboard[7])
    for row, expected in enumerate(EXPECTED_DASHBOARD, start=8):
        cells = dashboard[row]
        assert tuple(cell.value for cell in cells) == expected[:6]
        assert cells[2].number_format == analyzer.PERCENT_FORMAT
        assert cells[5].hyperlink.target == expected[6]

    # ---- Stock Data: every stock alphabetically by market then name, with formatted numbers
    data_sheet = workbook["Stock Data"]
    rows = list(data_sheet.iter_rows(values_only=True))
    assert rows[:len(STOCKS) + 1] == [tuple(analyzer.XLSX_HEADER)] + [xlsx_row(acronym) for acronym in XLSX_ORDER]
    assert data_sheet.cell(row=len(STOCKS) + 3, column=1).value == "Hover any header for a definition."
    assert data_sheet.freeze_panes == "B2"
    assert data_sheet.auto_filter.ref == "A1:G21"
    assert all(cell.comment is not None and cell.font.bold for cell in data_sheet[1])
    assert data_sheet["D2"].number_format == analyzer.PRICE_FORMAT
    assert data_sheet["F2"].number_format == analyzer.DIFFERENCE_FORMAT

    # each percent difference is marked ▲/▼ and shaded green (gain) or red (loss), darker the further from 0;
    # Apple's -25% is the biggest move, so it is the darkest; no change is a black 0 with no shading
    for row in range(2, len(STOCKS) + 2):
        cell = data_sheet.cell(row=row, column=7)
        assert cell.number_format == analyzer.PERCENT_FORMAT
        shade = analyzer.shade_for_rate(cell.value, 0.25)
        if cell.value == 0:
            assert shade is None and cell.fill.fill_type is None and cell.font.color.rgb[-6:] == "000000"
        else:
            assert cell.fill.fill_type == "solid" and cell.fill.start_color.rgb[-6:] == shade[0]
            assert cell.font.color.rgb[-6:] == shade[1]
    apple_row = XLSX_ORDER.index("AAPL:NASDAQ") + 2
    assert data_sheet.cell(row=apple_row, column=7).fill.start_color.rgb[-6:] == "B3261E"

    # if the AI failed, the table still lists the movers and their news links, without summaries
    analyzer.develop_xlsx(sorted_data, analyzer.AI_ERROR_ANSWER, xlsx_location)
    dashboard = openpyxl.load_workbook(xlsx_location)["Dashboard"]
    assert [cell.value for cell in dashboard[8]] == [COW_FACE + " Gainer", "UBER", 0.2, "—", "No summary available.",
                                                     "Uber Takes Robotaxis to Ten More Cities (Example News)"]
    assert dashboard[8][5].hyperlink.target == CANNED_NEWS["UBER"][2]

    # a stock without news says so, with no link
    analyzer.develop_xlsx(analyzer.sort_findings([finding(stock_id) for (stock_id, _, _, _, _) in STOCKS]),
                          CANNED_INSIGHT, xlsx_location)
    dashboard = openpyxl.load_workbook(xlsx_location)["Dashboard"]
    assert [cell.value for cell in dashboard[8]][3:] == ["Product / orders", "No recent news found.", "—"]
    assert dashboard[8][5].hyperlink is None
